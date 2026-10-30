"""Create an isolated, independently verified next candidate. No model calls."""
import json,sys,shutil,importlib.util,hashlib
from pathlib import Path
from openpyxl import load_workbook
project=Path(__file__).resolve().parents[2]
old=project/'runs/allocation-batch-v1-20261026'
root=project/'runs/allocation-diamond-v2-20261026'
root.mkdir(exist_ok=False)
shutil.copytree(old/'code',root/'code',ignore=shutil.ignore_patterns('__pycache__'))
sys.path.insert(0,str(root/'code'))
from tooluse_pilot.allocation_task import generate,reference,calculate
spec=importlib.util.spec_from_file_location('oracle',old/'verify_batch.py')
oracle=importlib.util.module_from_spec(spec);spec.loader.exec_module(oracle)
folder=root/'candidates/case-2';case=generate(1,folder)
w=load_workbook(folder/'reconciliation.xlsx')
data=[['A',137,17,1,1,'active','Red'],['B',211,13,1,1,'active','Blue'],['C',173,11,1,1,'active','Green'],['D',149,19,1,1,'active','Red'],['E',227,9,1,1,'active','Blue'],['F',193,14,1,1,'active','Green'],['G',167,16,1,1,'active','Red'],['H',239,8,1,1,'active','Blue'],['I',181,12,1,0,'active','Green'],['J',157,15,1,1,'active','Red'],['K',251,10,1,1,'active','Blue'],['L',199,13,1,1,'active','Green'],['M',263,18,1,1,'inactive','Red'],['N',229,11,0,1,'active','Blue']]
tables={'Orders':[list(next(w['Orders'].values))]+data[::2]+data[1::2],
 'Corrections':[['ID','Quantity'],['A',19],['C',13],['E',12],['G',18],['K',14]],
 'Priority':[['ID','Rank']]+[[r[0],200-i*5] for i,r in enumerate(data)],
 'Dependencies':[['ID','Requires'],['D','A'],['D','B'],['E','B'],['F','C'],['G','D'],['G','E'],['H','E'],['H','F'],['J','G'],['K','H'],['L','F'],['L','I']],
 'Limits':[['Group','Capacity'],['Red',8500],['Blue',8000],['Green',7500]],
 'Policy':[['Key','Value'],['Max per order',4500],['Cities','Beijing'],['Budget',19800]]}
for name,rows in tables.items():
 s=w[name];s.delete_rows(1,s.max_row)
 for row in rows:s.append(row)
w.save(folder/'reconciliation.xlsx')
initial=oracle.independent(w);assert initial==calculate(w);w.close()
case.update(id='allocation-diamond-v2-2',variant=2,revision={'Policy':{'Cities':'Beijing,Shenzhen','Budget':18500},'Corrections':{'A':23,'C':16,'E':15,'G':21,'K':17},'Limits':{'Red':7800,'Blue':9200,'Green':6700},'Priority':{'D':220,'F':215,'E':210,'H':205,'G':175,'I':202}})
(folder/'candidate.json').write_text(json.dumps(case,indent=2),encoding='utf-8')
ref=root/'references/case-2';shutil.copytree(folder,ref)
score,calls=reference(case,ref);assert score['success'] and len(calls)<=80
w=load_workbook(ref/'reconciliation.xlsx');revised=oracle.independent(w);assert revised==calculate(w);w.close();assert revised!=initial
record={'initial':{'summary':initial[0],'decisions':initial[1]},'revised':{'summary':revised[0],'decisions':revised[1]},'reference_score':score,'reference_calls':calls,'reference_tool_calls':len(calls),'model_calls':0}
(ref/'reference-result.json').write_text(json.dumps(record,indent=2),encoding='utf-8')
(root/'DESIGN.txt').write_text('Same explicit one-pass rules and tools. Add AND dependency diamonds across three capacity groups, multi-level propagation, non-round integer products, and revision that moves prerequisite-dependent orders ahead of parents. No new grader or hidden rules. Reference uses 28 tool calls; retain 25 actual API / 80 tool budget. Existing case1 recovered correct final answers, so it is not a confirmed hard case. Qwen screening and independent rerun required before user GPT comparison.\n',encoding='utf-8')
manifest={str(p.relative_to(root)):hashlib.sha256(p.read_bytes()).hexdigest() for p in root.rglob('*') if p.is_file() and '__pycache__' not in p.parts}
(root/'MANIFEST.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps({k:record[k] for k in ['initial','revised','reference_tool_calls','reference_score']},indent=2))
