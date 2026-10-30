import sys,json,shutil,hashlib
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent;sys.path.insert(0,str(ROOT/'code'))
from tooluse_pilot.w8_task import Environment
from independent_w7 import oracle
run=ROOT/sys.argv[1];result=json.loads((run/'result.json').read_text(encoding='utf-8'));artifact=run/result['artifact_dir'];out=run/(sys.argv[2] if len(sys.argv)>2 else 'audit');out.mkdir(exist_ok=False)
case=json.loads((artifact/'candidate.json').read_text(encoding='utf-8'));data=json.loads((artifact/'dataset.json').read_text(encoding='utf-8'))
replay=out/'replay';replay.mkdir()
for f in ('candidate.json','dataset.json'):shutil.copy2(artifact/f,replay/f)
e=Environment(case,replay);count=0
for event in result['events']:
 if event['event']!='tool_call':continue
 try:actual={'ok':True,'result':e.call(event['tool'],json.loads(event['arguments']))}
 except Exception as exc:e.invalid_actions+=1;actual={'ok':False,'error':str(exc)[:300]}
 assert actual==event['result'],(count,actual,event['result']);count+=1
assert e.score()==result['score']
wb=load_workbook(artifact/'submission.xlsx',data_only=True) if (artifact/'submission.xlsx').exists() else None
rows=[list(r) for r in wb['Top3'].iter_rows(min_row=2,values_only=True)] if wb else []
if wb:wb.close()
if case['condition']=='dynamic':
 for row in data['grades']:
  if row['student_id']=='s05' and row['batch_id']=='B2':row.update(record_id='g-B2-s05-r2',revision=2,score=96,supersedes='g-B2-s05-r1')
gold=oracle(data);independent=[r[:4] for r in rows]==[r[:4] for r in gold];assert independent==result['score']['answer_success']
wire=[json.loads(s) for s in (run/'wire-requests.jsonl').read_text().splitlines()];responses=[json.loads(s) for s in (run/'responses.jsonl').read_text().splitlines()]
assert all(w['url']=='http://10.10.23.7:5610/v1/chat/completions' and w['model']=='Qwen3.8-Flash-Next-AWQ-INT4' and w['chat_template_kwargs']=={'enable_thinking':False} for w in wire)
errors=sum(bool(x['error']) for x in result['events'] if x['event']=='api_attempt')
summary={'run':run.name,'answer_success':independent,'stop_reason':result.get('stop_reason'),'task_completed':result.get('task_completed'),'budget_exhausted':result.get('budget_exhausted'),'coverage_complete':result['score']['coverage_complete'],'evidence_valid':result['score']['evidence_valid'],'api_attempts':result['requests'],'wire_requests':len(wire),'successful_responses':len(responses),'api_errors':errors,'tools_replayed_exactly':count,'invalid_actions':result['invalid_tool_actions'],'expected':gold,'actual':rows,'response_models':sorted({r['model'] for r in responses}),'dataset_sha256':hashlib.sha256((artifact/'dataset.json').read_bytes()).hexdigest(),'candidate_sha256':hashlib.sha256((artifact/'candidate.json').read_bytes()).hexdigest()}
(out/'AUDIT.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(summary,ensure_ascii=False))
