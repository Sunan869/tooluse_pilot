"""A real XLSX workbook with two adjacent data regions and explicit table provenance."""
from pathlib import Path
from openpyxl import Workbook, load_workbook

def create_workbook(case, path):
    if case["family"] != "workbook_layout":
        raise ValueError("wrong family")
    path = Path(path)
    path.parent.mkdir(parents=True,exist_ok=True)
    env = case["environment"]
    row=int(env["target_cell"].split("C")[0][1:])
    col=int(env["target_cell"].split("C")[1])
    wb=Workbook()
    ws=wb.active; ws.title="Sales"
    ws.cell(row-2,col,"Current")
    ws.cell(row-2,col+1,"Archive")
    ws.cell(row-1,col,"Revenue")
    ws.cell(row-1,col+1,"Revenue")
    ws.cell(row,col,env["target_value"])
    ws.cell(row,col+1,env["archive_value"])
    ws.freeze_panes="A2"
    wb.save(path);wb.close()
    return path

def read_region(path, region):
    if region not in ("Current","Archive"):
        raise ValueError("Unknown region")
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        ws=wb["Sales"]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value==region:
                    c=cell.column
                    return {"region":region,"value":ws.cell(cell.row+2,c).value,
                            "source_cell":ws.cell(cell.row+2,c).coordinate}
        raise ValueError("Region not found")
    finally:
        wb.close()
