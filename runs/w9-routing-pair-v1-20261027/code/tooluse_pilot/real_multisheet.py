"""Actual multi-sheet workbook audit with duplicate IDs and inactive records."""
import json,random
from pathlib import Path
from openpyxl import Workbook,load_workbook

def generate(seed,folder):
    root=Path(folder);root.mkdir(parents=True,exist_ok=False)
    rng=random.Random(seed)
    rows=[("A-001",1200,"active",2),("A-002",1500,"active",1),
          ("A-001",1200,"active",1),("A-003",4000,"inactive",5),
          ("A-004",800,"active",3),("A-002",1500,"active",2)]
    rng.shuffle(rows)
    wb=Workbook();a=wb.active;a.title="Orders"
    a.append(["Order ID","Unit Price","Status","Quantity"])
    for row in rows:a.append(row)
    b=wb.create_sheet("Corrections");b.append(["Order ID","Correct Quantity"])
    b.append(["A-001",4]);b.append(["A-002",2])
    c=wb.create_sheet("Summary");c.append(["Metric","Value"])
    for name in ("Active total","Active unique orders","Highest value order"):c.append([name,None])
    wb.save(root/"audit.xlsx");wb.close()
    task={"id":f"multisheet-audit-{seed}","family":"real_multisheet_audit",
          "task":"Read Orders and Corrections in audit.xlsx. Group duplicate order IDs rather than adding duplicate rows. For corrected IDs use the corrected quantity; otherwise use the row quantity. Ignore inactive records. Compute total active order value, unique active order count, and highest-value active order ID (alphabetical tie break). Write to Summary B2:B4, save workbook and write report.txt with all three results; then finish.",
          "input":{"workbook":"audit.xlsx"}}
    (root/"candidate.json").write_text(json.dumps(task,indent=2),encoding="utf-8")
    return task

def oracle(book):
    corrections={r[0]:int(r[1]) for r in list(book["Corrections"].values)[1:]}
    orders={}
    for oid,price,status,qty in list(book["Orders"].values)[1:]:
        if status!="active":continue
        if oid in orders and orders[oid][0]!=int(price):raise ValueError("conflicting duplicate prices")
        orders[oid]=(int(price),corrections.get(oid,int(qty)))
    values={oid:price*qty for oid,(price,qty) in orders.items()}
    winner=sorted(values,key=lambda oid:(-values[oid],oid))[0]
    return (str(sum(values.values())),str(len(orders)),winner)

def specs():
    def t(n,description,properties=None,required=None):
        return {"type":"function","function":{"name":n,"description":description,
                "parameters":{"type":"object","properties":properties or {},"required":required or [],"additionalProperties":False}}}
    S={"type":"string"}
    return [t("list_sheets","List workbook sheets"),
            t("read_sheet","Read all rows in a named workbook sheet",{"sheet":S},["sheet"]),
            t("write_summary","Write output to one Summary cell",{"cell":S,"value":S},["cell","value"]),
            t("save_workbook","Persist workbook edits"),
            t("write_report","Write UTF-8 summary report",{"text":S},["text"]),
            t("finish","Finish task")]

class Environment:
    def __init__(self,case,root):
        self.root=Path(root);self.path=self.root/case["input"]["workbook"]
        self.wb=load_workbook(self.path);self.calls=[];self.finished=False;self.invalid_actions=0
    def call(self,name,args):
        if self.finished:raise ValueError("already finished")
        if name=="list_sheets" and args=={}:out={"sheets":self.wb.sheetnames}
        elif name=="read_sheet" and set(args)=={"sheet"} and args["sheet"] in self.wb.sheetnames:
            out={"rows":[list(r) for r in self.wb[args["sheet"]].values]}
        elif name=="write_summary" and set(args)=={"cell","value"} and args["cell"] in ("B2","B3","B4") and isinstance(args["value"],str):
            self.wb["Summary"][args["cell"]]=args["value"];out={"ok":True}
        elif name=="save_workbook" and args=={}:self.wb.save(self.path);out={"saved":True}
        elif name=="write_report" and set(args)=={"text"} and isinstance(args["text"],str):
            (self.root/"report.txt").write_text(args["text"],encoding="utf-8");out={"saved":True}
        elif name=="finish" and args=={}:self.finished=True;out={"ok":True}
        else:raise ValueError("invalid arguments")
        self.calls.append({"tool":name,"args":dict(args),"response":out})
        return out
    def score(self):
        book=load_workbook(self.path,read_only=True,data_only=True)
        try:
            correct=oracle(book)
            actual=tuple(str(book["Summary"][c].value or "") for c in ("B2","B3","B4"))
        finally:book.close()
        text=(self.root/"report.txt").read_text(encoding="utf-8") if (self.root/"report.txt").is_file() else ""
        checks={"summary":actual==correct,"report":all(x in text for x in correct),
                "finished":self.finished}
        return {"success":all(checks.values()) and not self.invalid_actions,
                "checks":checks,"expected":correct,"actual":actual,
                "tool_calls":len(self.calls),"invalid_actions":self.invalid_actions}
    def close(self):self.wb.close()

def reference(case,root):
    env=Environment(case,root)
    try:
        for name in ("Orders","Corrections"):env.call("read_sheet",{"sheet":name})
        answer=oracle(env.wb)
        for cell,val in zip(("B2","B3","B4"),answer):
            env.call("write_summary",{"cell":cell,"value":val})
        env.call("save_workbook",{})
        env.call("write_report",{"text":f"Total {answer[0]}; unique orders {answer[1]}; highest order {answer[2]}."})
        env.call("finish",{})
        return env.score()
    finally:env.close()
