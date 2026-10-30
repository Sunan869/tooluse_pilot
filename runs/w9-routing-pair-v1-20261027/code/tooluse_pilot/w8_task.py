import copy,json
from openpyxl import load_workbook
from .w7_task import Environment as Base,specs as base_specs
class Environment(Base):
 terminal_stops_runner=True
 def __init__(self,case,folder):
  super().__init__(case,folder);self.revision='R1';self.published_update=False;self.saved_revision=None;self.data['snapshot_id']='W8-R1'
 def publish(self):
  if self.case['condition']!='dynamic' or self.published_update:return
  for row in self.data['grades']:
   if row['record_id']=='g-B2-s05-r1':row.update(record_id='g-B2-s05-r2',revision=2,score=96,supersedes='g-B2-s05-r1')
  self.revision='R2';self.data['snapshot_id']='W8-R2';self.published_update=True
 def status(self):
  return {'current_revision':self.revision,'update_published':self.published_update,'changed_record_ids':['g-B2-s05-r2'] if self.published_update else [],'saved_revision':self.saved_revision,'recovery':'If read_revision differs, re-read authoritative sources; then write using current expected_revision. No further update after R2.'}
 def call(self,name,args):
  if self.finished:raise ValueError('Terminal task: runner must stop after successful finish')
  before=len(self.calls)
  if name=='get_status':result=self.status()
  elif name=='read_submission':
   if not (self.folder/'submission.xlsx').exists():result={'exists':False,'saved_revision':None,'rows':[]}
   else:
    wb=load_workbook(self.folder/'submission.xlsx',data_only=True);rows=[list(r) for r in wb['Top3'].iter_rows(min_row=2,values_only=True)];wb.close();result={'exists':True,'saved_revision':self.saved_revision,'rows':rows}
  elif name=='write_submission':
   self.publish()
   if args['expected_revision']!=self.revision:result={'accepted':False,'reason':'stale_revision','current_revision':self.revision,'saved_revision':self.saved_revision,'recovery':'Query get_status, re-read changed authoritative sources and resubmit. No file was modified.'}
   else:
    result=super().call(name,{'rows_json':args['rows_json']});self.saved_revision=self.revision;result.update(accepted=True,saved_revision=self.saved_revision)
  elif name=='finish':
   if self.saved_revision!=self.revision:result={'finished':False,'reason':'missing_or_stale_submission','current_revision':self.revision,'saved_revision':self.saved_revision,'recovery':'Save a submission with the current revision before finish.'}
   else:result=super().call(name,args)
  else:
   read_revision=self.revision;result=super().call(name,args)
   if name=='directory':result.update(version_rules='At most one publication follows first grades read, or precedes first write. Reads expose both revisions; query get_status before writing. write_submission compares expected_revision but never checks answer. Wrong saves can be replaced until finish. Successful finish is terminal; runner stops immediately.',current_revision=self.revision)
   if name in ('read_rows','get_record'):
    if args['source']=='grades':self.publish()
    result.update(read_revision=read_revision,current_revision=self.revision,stale=read_revision!=self.revision)
    if result['stale']:result['recovery']='Revision changed after this read. Query get_status and re-read the authoritative data before submission.'
  if len(self.calls)>before:self.calls.pop()
  self.calls.append({'tool':name,'arguments':copy.deepcopy(args),'result':copy.deepcopy(result)})
  return result
 def score(self):
  s=super().score();s.update(task_completed=self.finished,answer_correct=s['answer_success'],saved_revision=self.saved_revision,current_revision=self.revision,revision_current=self.saved_revision==self.revision)
  s['success']=s['answer_success'] and s['task_completed'] and s['revision_current'];return s

def specs():
 tools=base_specs()
 for tool in tools:
  f=tool['function']
  if f['name']=='write_submission':
   f['parameters']['properties']['expected_revision']={'type':'string'};f['parameters']['required'].append('expected_revision');f['description']+=' Requires expected_revision from current status; stale write has no effect and is recoverable. Can overwrite before finish.'
  if f['name']=='finish':f['description']='Irreversible terminal commit. Stops runner immediately. Check saved submission and version first. Does not check answer correctness. Missing/stale submission is recoverable and nonterminal.'
 for name,description in [('get_status','Read current revision, changed record IDs and saved revision. Does not return answers.'),('read_submission','Inspect saved XLSX rows and their revision without grading them.')]:tools.append({'type':'function','function':{'name':name,'description':description,'parameters':{'type':'object','properties':{},'required':[],'additionalProperties':False}}})
 return tools
