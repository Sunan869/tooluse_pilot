import json,base64,copy
from pathlib import Path
from openpyxl import Workbook,load_workbook
SOURCES=['enrollment','grades','cache']
HEADERS=['student_id','name','score','batch_id','grade_record_id','enrollment_record_id']
def answer_rows(data):
    batch=data['target_batch'];latest={}
    for r in data['grades']:
        if r['batch_id']==batch and r['published'] and r['valid']:
            if r['student_id'] not in latest or r['revision']>latest[r['student_id']]['revision']:latest[r['student_id']]=r
    rows=[]
    for e in data['enrollment']:
        if e['batch_id']==batch and e['status']=='active' and e['student_id'] in latest:
            g=latest[e['student_id']];rows.append([e['student_id'],e['name'],g['score'],batch,g['record_id'],e['record_id']])
    return sorted(rows,key=lambda r:(-r[2],r[0]))[:3]
class Environment:
    def __init__(self,case,folder):
        self.case=case;self.folder=Path(folder);self.data=json.loads((self.folder/case['input']['dataset']).read_text(encoding='utf-8'))
        self.calls=[];self.invalid_actions=0;self.finished=False;self.seen={s:set() for s in SOURCES}
    def source(self,s):
        if s not in SOURCES:raise ValueError('Unknown source')
        rows=copy.deepcopy(self.data[s])
        if s=='cache' and self.case['condition']=='conflict':
            for r in rows:
                if r['student_id'] in ('s02','s06'):r.update(score=99 if r['student_id']=='s02' else 89,as_of_revision=1,score_record_id='g-B2-'+r['student_id']+'-r1')
        return rows
    def call(self,name,args):
        if self.finished:raise ValueError('Already finished')
        if name=='directory':
            result={'snapshot_id':self.data['snapshot_id'],'target_batch':self.data['target_batch'],'rules':'enrollment determines eligibility and name. grades determines score: target batch, published=true, valid=true, highest revision per student_id. Superseded old versions remain in source. Join by student_id and batch_id, never name. Cache is advisory and may be stale. Rank score descending, student_id ascending, top 3.','sources':{s:{'total_rows':len(self.data[s]),'fields':list(self.data[s][0]),'authority':s!='cache'} for s in SOURCES},'pagination':'Follow next_cursor while has_more, keeping source and batch_id filter. get_record is an alternative. Snapshot is fixed.','submission_columns':HEADERS}
        elif name in ('read_rows','get_record'):
            s=args['source'];rows=self.source(s)
            if name=='get_record':
                rows=[r for r in rows if r['record_id']==args['record_id']];result={'found':bool(rows),'record':rows[0] if rows else None}
            else:
                batch=args.get('batch_id');rows=[r for r in rows if batch is None or r['batch_id']==batch];total=len(rows);offset=0
                if args.get('cursor'):
                    try:
                        token=json.loads(base64.urlsafe_b64decode(args['cursor']).decode());offset=token[3]
                        if token[:3]!=[self.data['snapshot_id'],s,batch] or type(offset)!=int or not 0<=offset<total:raise ValueError()
                    except Exception:raise ValueError('Invalid cursor for this snapshot/source/filter')
                size=4 if self.case['condition']=='coverage' else max(total,1);rows=rows[offset:offset+size];more=offset+len(rows)<total
                cursor=base64.urlsafe_b64encode(json.dumps([self.data['snapshot_id'],s,batch,offset+len(rows)]).encode()).decode() if more else None
                result={'rows':rows,'total_matching':total,'has_more':more,'next_cursor':cursor}
            self.seen[s].update(r['record_id'] for r in rows);result.update(source=s,authority=s!='cache',snapshot_id=self.data['snapshot_id'])
        elif name=='write_submission':
            rows=json.loads(args['rows_json'])
            if not isinstance(rows,list) or any(not isinstance(r,list) or len(r)!=6 or type(r[2]) not in (int,float) for r in rows):raise ValueError('rows_json needs six-column rows with numeric score')
            wb=Workbook();ws=wb.active;ws.title='Top3';ws.append(HEADERS)
            for r in rows:ws.append(r)
            wb.save(self.folder/'submission.xlsx');wb.close();result={'saved':'submission.xlsx','rows':len(rows)}
        elif name=='finish':
            if not (self.folder/'submission.xlsx').is_file():raise ValueError('Save submission first')
            self.finished=True;result={'finished':True}
        else:raise ValueError('Unknown tool')
        self.calls.append({'tool':name,'arguments':copy.deepcopy(args),'result':copy.deepcopy(result)});return result
    def score(self):
        actual=[];expected=answer_rows(self.data)
        if (self.folder/'submission.xlsx').exists():
            wb=load_workbook(self.folder/'submission.xlsx',data_only=True);actual=[list(r) for r in wb['Top3'].iter_rows(min_row=2,values_only=True)];wb.close()
        answer=[r[:4] for r in actual]==[r[:4] for r in expected]
        coverage={s:{'seen':len(self.seen[s]&{r['record_id'] for r in self.data[s] if r['batch_id']==self.data['target_batch']}),'required':sum(r['batch_id']==self.data['target_batch'] for r in self.data[s])} for s in ('enrollment','grades')}
        complete=all(c['seen']==c['required'] for c in coverage.values());evidence=actual==expected and all(r[4] in self.seen['grades'] and r[5] in self.seen['enrollment'] for r in actual);success=answer and self.finished
        return {'success':success,'answer_success':answer,'finished':self.finished,'evidence_valid':evidence,'coverage_complete':complete,'coverage':coverage,'strict_process_success':success and self.invalid_actions==0,'invalid_actions':self.invalid_actions,'expected':expected,'actual':actual}
    def close(self):pass
def specs():
    source={'type':'string','enum':SOURCES}
    definitions=[('directory','Read authority, schemas, snapshot and task rules.',{},[]),('read_rows','Read source rows. Follow next_cursor with same filter until has_more=false.',{'source':source,'batch_id':{'type':'string'},'cursor':{'type':'string'}},['source']),('get_record','Read one source record by unique record_id.',{'source':source,'record_id':{'type':'string'}},['source','record_id']),('write_submission','Save Top3 XLSX. rows_json encodes rows [student_id,name,score,batch_id,grade_record_id,enrollment_record_id].',{'rows_json':{'type':'string'}},['rows_json']),('finish','Finish after saving. Does not validate correctness.',{},[])]
    return [{'type':'function','function':{'name':n,'description':d,'parameters':{'type':'object','properties':p,'required':r,'additionalProperties':False}}} for n,d,p,r in definitions]
