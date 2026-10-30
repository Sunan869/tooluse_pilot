"""Offline independent oracle + real tool reference + stale-output mutations."""
import json,sys,shutil
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'code'))
from tooluse_pilot.allocation_task import generate,reference,calculate,Environment

def independent(w):
    tables={s:[dict(zip(next(w[s].values),row)) for row in list(w[s].values)[1:]] for s in ['Orders','Corrections','Priority','Dependencies']}
    policy=dict(list(w['Policy'].values)[1:]);limits=dict(list(w['Limits'].values)[1:])
    selected=[];decisions=[]
    rank={x['ID']:x['Rank'] for x in tables['Priority']}
    corrections={x['ID']:x['Quantity'] for x in tables['Corrections']}
    for order in sorted(tables['Orders'],key=lambda x:(-rank[x['ID']],x['ID'])):
        n=order['ID'];amount=order['Price']*corrections.get(n,order['Qty'])
        parents=[x['Requires'] for x in tables['Dependencies'] if x['ID']==n]
        tests=[('inactive',order['Status']!='active'),('city',any(order[c]!=1 for c in policy['Cities'].split(','))),('per_order',amount>policy['Max per order']),('dependency',not set(parents).issubset(x[0] for x in selected)),('group_cap',sum(x[1] for x in selected if x[2]==order['Group'])+amount>limits[order['Group']]),('budget',sum(x[1] for x in selected)+amount>policy['Budget'])]
        reason=next((reason for reason,bad in tests if bad),'selected')
        if reason=='selected':selected.append((n,amount,order['Group']))
        decisions.append([n,amount,int(reason=='selected'),reason])
    best=sorted(selected,key=lambda x:(-x[1],x[0]))[0] if selected else ('NONE',0,None)
    return list(map(str,[sum(x[1] for x in selected),len(selected),best[0],best[1]])),decisions

def main():
    output=[]
    for v in [0,1]:
        folder=ROOT/'candidates'/f'case-{v}';case=generate(v,folder)
        wb=load_workbook(folder/'reconciliation.xlsx');initial=independent(wb);assert initial==calculate(wb);wb.close()
        ref=ROOT/'references'/f'case-{v}';shutil.copytree(folder,ref);score,calls=reference(case,ref)
        assert score['success'] and len(calls)<=80
        wb=load_workbook(ref/'reconciliation.xlsx');final=independent(wb);assert final==calculate(wb)
        assert final!=initial
        # Boundary/ordering probes on an in-memory copy, never model input files.
        probes=[]
        for budget in [0,4800,8000,14500,100000]:
            wb['Policy']['B4']=budget;assert independent(wb)==calculate(wb);probes.append(budget)
        wb.close()
        # Reconstruct legitimate terminal state, then introduce stale output only.
        mutation=ROOT/'mutations'/f'case-{v}';shutil.copytree(folder,mutation)
        env=Environment(case,mutation)
        for call in calls:env.call(call['tool'],call['args'])
        good=env.score();assert good['success']
        ws=env.wb['Decisions'];ws.delete_rows(2,ws.max_row)
        for row in initial[1]:ws.append(row)
        env.wb.save(env.path);stale=env.score();assert not stale['checks']['decisions'] and not stale['success']
        env.close()
        record={'case':case['id'],'initial':{'summary':initial[0],'decisions':initial[1]},'revised':{'summary':final[0],'decisions':final[1]},'reference_score':score,'reference_tool_calls':len(calls),'reference_calls':calls,'budget_boundary_probes':probes,'stale_decisions_rejected':True}
        (ref/'reference-result.json').write_text(json.dumps(record,indent=2));output.append(record)
    (ROOT/'VALIDATION.json').write_text(json.dumps({'model_calls':0,'cases':output},indent=2))
    print(json.dumps([{k:r[k] for k in ['case','reference_tool_calls','stale_decisions_rejected']} for r in output],indent=2))

if __name__=='__main__':
    if len(sys.argv)>1:
        ROOT=Path(sys.argv[1]).resolve();ROOT.mkdir(parents=True,exist_ok=False)
    main()
