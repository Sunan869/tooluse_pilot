"""Private-oracle deterministic scorer. Never expose this module to agents."""

import csv, json
from pathlib import Path
from openpyxl import load_workbook


def grade(case_dir, trace_path=None, trace_data=None):
    root = Path(case_dir)
    revision=root/'private'/'revision_applied.json'
    applied=revision.exists() and json.loads(revision.read_text(encoding='utf-8')).get('applied') is True
    oracle_name='oracle_revised.json' if applied else 'oracle.json'
    truth=json.loads((root/'private'/oracle_name).read_text(encoding='utf-8'))['rows']
    target = root / "output" / "anomalies.csv"
    if not target.exists():
        return {
            "total": 0,
            "result_score": 0,
            "process_score": 0,
            "error": "output missing",
        }
    try:
        with target.open(encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            if reader.fieldnames != [
                "order_id",
                "expected_cents",
                "net_paid_cents",
                "issues",
            ]:
                raise ValueError("CSV columns incorrect")
            raw = list(reader)
        rows = [
            {
                "order_id": r["order_id"],
                "expected_cents": int(r["expected_cents"]),
                "net_paid_cents": int(r["net_paid_cents"]),
                "issues": r["issues"],
            }
            for r in raw
        ]
        expected = {r["order_id"]: r for r in truth}
        actual = {r["order_id"]: r for r in rows}
        duplicates = len(actual) != len(rows)
        correct = sum(1 for key, item in expected.items() if actual.get(key) == item)
        extras = sum(1 for key in actual if key not in expected)
        result_score = max(0, correct - extras) / max(1, len(expected))
        if duplicates or rows != sorted(rows, key=lambda r: r["order_id"]):
            result_score = min(result_score, 0.8)
    except (ValueError, KeyError, TypeError, UnicodeError, csv.Error) as ex:
        return {"total": 0, "result_score": 0, "process_score": 0, "error": str(ex)}
    # Process is optional to support uninstrumented Python/Shell execution; not awarded without evidence.
    process_score = 0
    if trace_path or trace_data is not None:
        trace = trace_data if trace_data is not None else json.loads(Path(trace_path).read_text(encoding='utf-8'))
        names = [x.get("tool") for x in trace.get("calls", [])]
        if {"list_sheets", "read_rows", "write_csv"} <= set(names):
            sheetset = {
                x.get("args", {}).get("sheet")
                for x in trace.get("calls", [])
                if x.get("tool") == "read_rows"
            }

            wb = load_workbook(root / "input" / "procurement.xlsx", read_only=True)

            try:
                required = set(wb.sheetnames)
            finally:
                wb.close()
            if required <= sheetset:
                process_score = 1
    return {
        "total": round(100 * (0.85 * result_score + 0.15 * process_score), 2),
        "result_score": round(result_score, 4),
        "process_score": process_score,
        "expected_rows": len(expected),
        "actual_rows": len(rows),
        "correct_rows": correct,
        "extra_rows": extras,
        "duplicate_ids": duplicates,
    }

def delta_score(case_dir):
    """Validate changed order IDs against the revised oracle and actual output.

    A missing, extra, or partially updated changed record is counted incorrect.
    Unchanged entries are assessed by the ordinary result score.
    """
    root = Path(case_dir)
    before = json.loads((root / "private" / "oracle.json").read_text(encoding="utf-8"))["rows"]
    after = json.loads((root / "private" / "oracle_revised.json").read_text(encoding="utf-8"))["rows"]
    left = {r["order_id"]: r for r in before}
    right = {r["order_id"]: r for r in after}
    changed = {k for k in (left.keys() | right.keys()) if left.get(k) != right.get(k)}
    output = root / "output" / "anomalies.csv"
    actual = {}
    if output.is_file():
        try:
            with output.open("r", encoding="utf-8-sig", newline="") as f:
                raw = list(csv.DictReader(f))
            if len({r["order_id"] for r in raw}) != len(raw):
                return {"changed_count": len(changed), "changed_correct": 0, "changed_accuracy": 0.0}
            actual = {
                r["order_id"]: {
                    "order_id": r["order_id"], "expected_cents": int(r["expected_cents"]),
                    "net_paid_cents": int(r["net_paid_cents"]), "issues": r["issues"],
                } for r in raw
            }
        except (OSError, ValueError, KeyError, csv.Error):
            pass
    hits = sum(actual.get(k) == right.get(k) for k in changed)
    return {
        "changed_count": len(changed),
        "changed_correct": hits,
        "changed_accuracy": round(hits / len(changed), 4) if changed else 1.0,
    }
