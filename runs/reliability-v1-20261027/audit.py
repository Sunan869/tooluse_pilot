import sys,json,hashlib,shutil,copy
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent;sys.path.insert(0,str(ROOT/'code'))
from tooluse_pilot.w7_task import Environment as E7
from tooluse_pilot.w8_task import Environment as E8
from tooluse_pilot.w9_task import Environment as E9
from independent_w7 import oracle as sql_oracle
from independent_graph import oracle as graph_oracle

def audit(run):
 run=Path(run);raw=run/'result.json';result=json.loads(raw.read_text(encoding='utf-8'));digest=hashlib.sha256(raw.read_bytes()).hexdigest();out=run/'audit'
 if (out/'AUDIT.json').exists():
  previous=json.loads((out/'AUDIT.json').read_text());assert previous['result_sha256']==digest;return previous
 out.mkdir(exist_ok=False);artifact=run/result['artifact_dir'];case=json.loads((artifact/'candidate.json').read_text(encoding='utf-8'));data=json.loads((artifact/'dataset.json').read_text(encoding='utf-8'));fam=int(case['family'][-1]);replay=out/'replay';replay.mkdir()
 for n in ('candidate.json','dataset.json'):shutil.copy2(artifact/n,replay/n)
 env={7:E7,8:E8,9:E9}[fam](case,replay);count=0
 for event in result['events']:
  if event['event']!='tool_call':continue
  try:value={'ok':True,'result':env.call(event['tool'],json.loads(event['arguments']))}
  except Exception as exc:env.invalid_actions+=1;value={'ok':False,'error':str(exc)[:300]}
  assert value==event['result'],(count,value,event['result']);count+=1
 assert env.score()==result['score']
 if fam in (7,8):
  final=copy.deepcopy(data)
  if fam==8 and case['condition']=='dynamic':
   row=next(g for g in final['grades'] if g['student_id']=='s05' and g['batch_id']=='B2');row.update(record_id='g-B2-s05-r2',revision=2,score=96,supersedes='g-B2-s05-r1')
  gold=sql_oracle(final);actual=[]
  if (artifact/'submission.xlsx').exists():
   wb=load_workbook(artifact/'submission.xlsx',data_only=True);actual=[list(r) for r in wb['Top3'].iter_rows(min_row=2,values_only=True)];wb.close()
  correct=[r[:4] for r in actual]==[r[:4] for r in gold];assert correct==result['score']['answer_success']
 else:
  gold=graph_oracle(data,case['condition']);state=json.loads((artifact/'state.json').read_text()) if (artifact/'state.json').exists() else {'charges':[],'package':{'total_cost':0}}
  loc=data['start'];time=data['start_time'];cost=0;ids=[]
  for e in state['charges']:
   leg=next(x for x in data['routes'] if x['route_id']==e['route_id']);assert leg['origin']==loc and leg['departure']>=time and e['amount']==leg['cost'];assert case['condition']!='permanent' or leg['route_id']!=data['fault_route'];loc=leg['destination'];time=leg['arrival'];cost+=leg['cost'];ids.append(leg['route_id'])
  checks={'arrived':loc==data['destination'],'deadline':time<=data['deadline'],'budget':cost<=data['budget'],'minimum_cost':cost==gold['minimum_cost'],'no_duplicate_charges':len(ids)==len(set(ids)) and cost==state['package']['total_cost']};assert checks==result['score']['checks'];correct=all(checks.values());actual={'path':ids,'cost':cost}
 attempts=[e for e in result['events'] if e['event']=='api_attempt'];errors=sum(e['error'] is not None for e in attempts);wire=[json.loads(s) for s in (run/'wire-requests.jsonl').read_text().splitlines()] if (run/'wire-requests.jsonl').exists() else []
 assert all(w['model']=='Qwen3.8-Flash-Next-AWQ-INT4' and w['url']=='http://10.10.23.7:5610/v1/chat/completions' and w['chat_template_kwargs']=={'enable_thinking':False} for w in wire)
 labels=[];score=result['score'];events=[e for e in result['events'] if e['event']=='tool_call'];rejected=[e for e in events if not e['result']['ok'] or e['result'].get('result',{}).get('accepted') is False or e['result'].get('result',{}).get('success') is False or e['result'].get('result',{}).get('finished') is False]
 if errors:labels.append('API_failure')
 if result['state'] in ('max_requests','max_tools'):labels.append('budget_exhaustion')
 if result['invalid_tool_actions']:labels.append('invalid_tool_action')
 if not correct:
  if fam in (7,8):
   if actual:labels.append('selection_or_order')
   else:labels.append('missing_submission')
   if not score['coverage_complete']:labels.append('incomplete_coverage')
   if fam==8 and not score['revision_current']:labels.append('stale_version')
  else:
   if not score['checks']['minimum_cost']:labels.append('nonoptimal_cost')
   if not score['checks']['arrived']:labels.append('route_planning_or_early_stop')
   if any(e['code']=='MISSED_DEPARTURE' for e in score['rejections']):labels.append('missed_departure')
 if not result.get('task_completed',env.finished):labels.append('not_completed')
 if not correct and not labels:labels=['unknown']
 summary={'run':run.name,'family':fam,'condition':case['condition'],'case_id':case['id'],'business_correct':correct,'task_completed':bool(env.finished),'success':score['success'],'budget_exhausted':result['state'] in ('max_requests','max_tools'),'requests':result['requests'],'api_errors':errors,'tools':count,'invalid_actions':result['invalid_tool_actions'],'recoverable_rejections':len(rejected),'recovered_after_rejection_success':bool(rejected) and score['success'],'fault_exposed':score.get('fault_exposed'),'coverage_complete':score.get('coverage_complete'),'revision_current':score.get('revision_current'),'labels_observable_not_causal':labels,'gold':gold,'actual':actual,'result_sha256':digest,'candidate_sha256':hashlib.sha256((artifact/'candidate.json').read_bytes()).hexdigest(),'dataset_sha256':hashlib.sha256((artifact/'dataset.json').read_bytes()).hexdigest()}
 (out/'AUDIT.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8');return summary
if __name__=='__main__':print(json.dumps(audit(ROOT/sys.argv[1]),ensure_ascii=False))
