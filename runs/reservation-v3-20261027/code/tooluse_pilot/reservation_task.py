"""Three-round commitments with cascading release and locked amounts."""
import json
from openpyxl import load_workbook
from . import real_w6_reconciliation as base
from .allocation_task import specs as old_specs

REASONS={'retained','cancelled','inactive','city','per_order','dependency','group_cap','budget','selected'}

def calculate(w):
    orders={r[0]:r for r in list(w['Orders'].values)[1:]}
    qty=dict(list(w['Corrections'].values)[1:]);rank=dict(list(w['Priority'].values)[1:])
    policy=dict(list(w['Policy'].values)[1:]);caps=dict(list(w['Limits'].values)[1:])
    parents={n:[] for n in orders}
    for child,parent in list(w['Dependencies'].values)[1:]:parents[child].append(parent)
    withdrawn={r[0] for r in list(w['Withdrawn'].values)[1:]}
    held=dict(list(w['Commitments'].values)[1:]);held={n:a for n,a in held.items() if n not in withdrawn}
    while True:
        remove=[n for n in held if any(p not in held for p in parents[n])]
        if not remove:break
        for n in remove:del held[n]
    chosen=dict(held);used={g:0 for g in caps}
    for n,a in held.items():used[orders[n][6]]+=a
    total=sum(held.values());rows=[]
    for n in sorted(orders,key=lambda n:(-rank[n],n)):
        _,price,q,bj,sz,status,group=orders[n];amount=held.get(n,price*qty.get(n,q))
        if n in held:reason='retained'
        elif n in withdrawn:reason='cancelled'
        elif status!='active':reason='inactive'
        elif any({'Beijing':bj,'Shenzhen':sz}[c]!=1 for c in policy['Cities'].split(',')):reason='city'
        elif amount>policy['Max per order']:reason='per_order'
        elif any(p not in chosen for p in parents[n]):reason='dependency'
        elif used[group]+amount>caps[group]:reason='group_cap'
        elif total+amount>policy['Budget']:reason='budget'
        else:
            reason='selected';chosen[n]=amount;used[group]+=amount;total+=amount
        rows.append([n,amount,int(reason in {'selected','retained'}),reason])
    winner=min(chosen,key=lambda n:(-chosen[n],n)) if chosen else 'NONE'
    return list(map(str,[total,len(chosen),winner,chosen.get(winner,0)])),rows

def specs():
    result=old_specs()
    for s in result:
        f=s['function']
        if f['name']=='write_decisions':f['description']='Replace all Decisions rows in processing order. rows_json: [ID, integer amount, selected 0/1, reason]. Reasons: retained, cancelled, inactive, city, per_order, dependency, group_cap, budget, selected. Save afterwards.'
        if f['name']=='receive_update':f['description']='After saving correct current-round delivery, commit selected rows and receive the next update. Required twice: round 0 to 1, then 1 to 2. No update after round 2.'
        if f['name']=='finish':f['description']='Finish after saving round 2 delivery. Does not calculate or correct your outputs.'
    return result

class Environment(base.Environment):
    def __init__(self,case,folder):
        super().__init__(case,folder);self.round=0;self.passed=[];self.saved_round=False
    def disk_checks(self):
        w=load_workbook(self.path,read_only=True,data_only=True)
        try:
            expected,rows=calculate(w);actual=[str(w['Summary'][c].value if w['Summary'][c].value is not None else '') for c in ['B2','B3','B4','B5']]
            decisions=[list(r) for r in list(w['Decisions'].values)[1:]]
        finally:w.close()
        p=self.root/'report.txt';text=p.read_text(encoding='utf-8') if p.exists() else ''
        return expected,actual,{'result':expected==actual,'decisions':rows==decisions,'report':base.report_valid(text,expected)}
    def call(self,name,args):
        if self.finished:raise ValueError('already finished')
        if name=='write_decisions' and set(args)=={'rows_json'}:
            rows=json.loads(args['rows_json']);ids={r[0] for r in list(self.wb['Orders'].values)[1:]}
            if not isinstance(rows,list) or len(rows)!=len(ids):raise ValueError('one decision per order required')
            if any(not isinstance(r,list) or len(r)!=4 or r[0] not in ids or type(r[1]) is not int or type(r[2]) is not int or r[2] not in [0,1] or r[3] not in REASONS for r in rows):raise ValueError('invalid decision row')
            if len({r[0] for r in rows})!=len(ids):raise ValueError('duplicate order ID')
            s=self.wb['Decisions'];s.delete_rows(2,s.max_row)
            for row in rows:s.append(row)
            out={'written':len(rows)}
        elif name=='receive_update' and args=={}:
            if self.round>=2:raise ValueError('no more updates; finish round 2')
            _,_,checks=self.disk_checks();issues=[k for k,v in checks.items() if not v]
            if not self.saved_round:issues.append('workbook_not_saved')
            if issues:raise ValueError('round_'+str(self.round)+'_failed: '+', '.join(issues))
            self.wb.close();self.wb=load_workbook(self.path)
            commitments=[[r[0],r[1]] for r in list(self.wb['Decisions'].values)[1:] if r[2]==1]
            s=self.wb['Commitments'];s.delete_rows(2,s.max_row)
            for row in commitments:s.append(row)
            patch=self.case['updates'][self.round]
            for name2,changes in patch['tables'].items():
                sh=self.wb[name2];positions={str(sh.cell(i,1).value):i for i in range(2,sh.max_row+1)}
                for key,value in changes.items():
                    if key in positions:sh.cell(positions[key],2,value)
                    else:sh.append([key,value])
            known={r[0] for r in list(self.wb['Withdrawn'].values)[1:]}
            for n in patch['withdraw']:
                if n not in known:self.wb['Withdrawn'].append([n])
            self.passed.append(self.round);self.round+=1;self.wb['Stage']['B2']=self.round
            self.transitioned=True;self.initial_ok=True;self.saved_round=False;self.saved_after=False;self.wb.save(self.path)
            out={'round':self.round,'changed_sheets':['Commitments','Withdrawn','Stage',*patch['tables']],'instruction':'Keep surviving commitments at locked amounts, release cancelled/dependent commitments recursively, reserve survivors before processing new orders. Read changed sheets and save complete delivery.'}
        elif name=='finish' and args=={}:
            if self.round!=2:raise ValueError('finish requires round 2; receive two updates')
            self.finished=True;out={'ok':True}
        else:
            out=super().call(name,args)
            if name=='save_workbook':self.saved_round=True
            return out
        self.calls.append({'tool':name,'args':dict(args),'response':out});return out
    def score(self):
        expected,actual,checks=self.disk_checks()
        checks.update(round0_delivered=0 in self.passed,round1_delivered=1 in self.passed,round2_received=self.round==2,final_saved=self.saved_round and self.round==2,finished=self.finished)
        return {'success':all(checks.values()) and not self.invalid_actions,'checks':checks,'expected':expected,'actual':actual,'invalid_actions':self.invalid_actions,'tool_calls':len(self.calls),'round':self.round}
