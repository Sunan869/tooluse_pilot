"""Private-oracle deterministic scorer. Never expose this module to agents."""

import csv, json
from pathlib import Path
from openpyxl import load_workbook


def grade(case_dir, trace_path=None):
    root = Path(case_dir)
    truth = json.loads((root / "private" / "oracle.json").read_text(encoding="utf-8"))[
        "rows"
    ]
    target = root / "output" / "anomalies.csv"
    if not target.exists():
        return {
            "total": 0,
            "result_score": 0,
            "process_score": 0,
            "error": "output missing",
        }
    try:
        with target.open(encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            if reader.fieldnames != [
                "order_id",
                "expected_cents",
                "net_paid_cents",
                "issues",
            ]:
                raise ValueError("CSV columns incorrect")
            raw = list(reader)
        rows = [
            {
                "order_id": r["order_id"],
                "expected_cents": int(r["expected_cents"]),
                "net_paid_cents": int(r["net_paid_cents"]),
                "issues": r["issues"],
            }
            for r in raw
        ]
        expected = {r["order_id"]: r for r in truth}
        actual = {r["order_id"]: r for r in rows}
        duplicates = len(actual) != len(rows)
        correct = sum(1 for key, item in expected.items() if actual.get(key) == item)
        extras = sum(1 for key in actual if key not in expected)
        result_score = max(0, correct - extras) / max(1, len(expected))
        if duplicates or rows != sorted(rows, key=lambda r: r["order_id"]):
            result_score = min(result_score, 0.8)
    except (ValueError, KeyError, TypeError, UnicodeError, csv.Error) as ex:
        return {"total": 0, "result_score": 0, "process_score": 0, "error": str(ex)}
    # Process is optional to support uninstrumented Python/Shell execution; not awarded without evidence.
    process_score = 0
    if trace_path:
        trace = json.loads(Path(trace_path).read_text(encoding="utf-8"))
        names = [x.get("tool") for x in trace.get("calls", [])]
        if {"list_sheets", "read_rows", "write_csv"} <= set(names):
            sheetset = {
                x.get("args", {}).get("sheet")
                for x in trace.get("calls", [])
                if x.get("tool") == "read_rows"
            }

            wb = load_workbook(root / "input" / "procurement.xlsx", read_only=True)

            try:
                required = set(wb.sheetnames)
            finally:
                wb.close()
            if required <= sheetset:
                process_score = 1
    return {
        "total": round(100 * (0.85 * result_score + 0.15 * process_score), 2),
        "result_score": round(result_score, 4),
        "process_score": process_score,
        "expected_rows": len(expected),
        "actual_rows": len(rows),
        "correct_rows": correct,
        "extra_rows": extras,
        "duplicate_ids": duplicates,
    }
