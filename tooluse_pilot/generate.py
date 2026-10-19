"""Deterministic workbook + secret expected answer generation."""
import json, random
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter
from .core import compute

HEADERS={
'Orders':['order_id','supplier_id','dept','amount_cents','status','test_flag'],
'Payments':['payment_id','order_id','kind','amount_cents'],
'Suppliers':['supplier_id','enabled'],
'Rules':['dept','approval_limit_cents'],
}

def make_case(root,seed=20261019,count=40):
    root=Path(root); inp=root/'input'; private=root/'private'; inp.mkdir(parents=True,exist_ok=True); private.mkdir(parents=True,exist_ok=True)
    rng=random.Random(seed)
    data={'Orders':[],'Payments':[],'Suppliers':[], 'Rules':[]}
    data['Suppliers']=[{'supplier_id':f'S{i:02d}','enabled':int(i!=4)} for i in range(1,9)]
    data['Rules']=[{'dept':'ENG','approval_limit_cents':2000000},{'dept':'OPS','approval_limit_cents':900000}]
    payment_seq=0
    for i in range(count):
        oid=f'PO-{i+1:04d}'; sid=rng.choice([f'S{j:02d}' for j in range(1,9)])
        amount=rng.randint(20000,3000000)
        status='CANCELLED' if i%11==0 else 'ACTIVE'; test=int(i%13==0)
        data['Orders'].append({'order_id':oid,'supplier_id':sid,'dept':rng.choice(['ENG','OPS']),'amount_cents':amount,'status':status,'test_flag':test})
        paid=amount if i%7!=0 else amount-1000
        payment_seq+=1
        data['Payments'].append({'payment_id':f'PAY-{payment_seq:04d}','order_id':oid,'kind':'PAYMENT','amount_cents':paid})
        if i%5==0:
            payment_seq+=1
            data['Payments'].append({'payment_id':f'PAY-{payment_seq:04d}','order_id':oid,'kind':'REFUND','amount_cents':rng.randint(100,3000)})
    wb=Workbook()
    wb.remove(wb.active)
    for name,rows in data.items():
        sheet=wb.create_sheet(name); headers=HEADERS[name]
        sheet.append(headers)
        for r in rows: sheet.append([r[h] for h in headers])
        for cell in sheet[1]:
            cell.fill=PatternFill(fill_type='solid',fgColor='17365D')
            cell.font=Font(bold=True,color='FFFFFF')
        sheet.row_dimensions[1].height=26
        for c in range(1,len(headers)+1): sheet.column_dimensions[get_column_letter(c)].width=22
        sheet.freeze_panes='A2'
    wb.save(inp/'procurement.xlsx')
    oracle=compute(data)
    (private/'oracle.json').write_text(json.dumps({'schema_version':'1.0','rows':oracle},ensure_ascii=False,indent=2),encoding='utf-8')
    (private/'generator_config.json').write_text(json.dumps({'seed':seed,'count':count},indent=2),encoding='utf-8')
    instruction='''# 采购异常核验任务\n\n读取 input/procurement.xlsx 中 Orders、Payments、Suppliers、Rules 四个工作表。\n对 status=ACTIVE 且 test_flag=0 的订单逐条核验：\n1. 供应商必须在 Suppliers 中且 enabled=1，否则标记 SUPPLIER_INVALID；\n2. 订单金额超过所属部门 approval_limit_cents 时标记 REQUIRES_APPROVAL；部门规则缺失标记 RULE_MISSING；\n3. 同一订单的 PAYMENT 金额合计减去 REFUND 金额合计，若不等于订单 amount_cents，标记 PAYMENT_MISMATCH。\n\n输出 output/anomalies.csv，CSV 列顺序为 order_id,expected_cents,net_paid_cents,issues。只输出存在异常的订单，issues 使用字典序排序后以 | 分隔，按 order_id 升序排列。金额单位全部为整数分。不要修改输入文件。\n'''
    (root/'instruction.md').write_text(instruction,encoding='utf-8')
    task={'schema_version':'1.0','task_id':'TOOL-XLSX-W07-W12-001','title':'采购订单跨表核验','primary_weakness':'W7','secondary_weaknesses':['W12','W4'],'allowed_modes':['structured','python_shell'],'input_files':['input/procurement.xlsx'],'output_file':'output/anomalies.csv','oracle_path':'private/oracle.json','tool_names':['list_sheets','read_rows','write_csv'],'grading':{'result_weight':0.85,'process_weight':0.15}}
    (root/'task.json').write_text(json.dumps(task,ensure_ascii=False,indent=2),encoding='utf-8')
    return root
