"""Multi-sheet, two-stage W6 XLSX reconciliation benchmark."""
import json,random,re
from pathlib import Path
from openpyxl import Workbook,load_workbook

def generate(seed,folder):
    root=Path(folder);root.mkdir(parents=True,exist_ok=False)
    wb=Workbook();orders=wb.active;orders.title="Orders"
    orders.append(["ID","Price","Qty","Beijing","Shenzhen","Status"])
    rows=[("Alpha",1000,8,1,0,"active"),("Beta",1800,3,1,1,"active"),("Gamma",2200,2,1,1,"active"),("Delta",2600,2,1,1,"active"),("Dormant",9000,10,1,1,"inactive")]
    random.Random(seed).shuffle(rows)
    for row in rows:orders.append(row)
    corr=wb.create_sheet("Corrections");corr.append(["ID","Quantity"])
    for name,qty in (("Beta",3),("Gamma",2),("Delta",2)):corr.append([name,qty])
    pol=wb.create_sheet("Policy")
    for row in (("Key","Value"),("Max per order",10000),("Cities","Beijing")):pol.append(row)
    summ=wb.create_sheet("Summary")
    for row in (("Metric","Value"),("Total",None),("Count",None),("Winner",None),("Winner amount",None)):summ.append(row)
    wb.save(root/"reconciliation.xlsx");wb.close()
    case={"id":f"w6-reconcile-{seed}","family":"real_w6_reconciliation","scenario":seed%3,
          "input":{"workbook":"reconciliation.xlsx"},
          "task":"Read Orders, Corrections, Policy. Active orders only; correction quantity replaces original. Filter by all required cities and per-order maximum. Save Summary B2:B5 as sum of eligible values, count, highest-value order ID, highest order amount (ties alphabetical). Write report.txt explicitly containing total, count, winner, amount. Then call receive_update; Policy and Corrections change. Re-read them and revise ALL outputs, save again, finish."}
    (root/"candidate.json").write_text(json.dumps(case,indent=2),encoding="utf-8")
    return case

def oracle(wb):
    corr={str(n):int(q) for n,q in list(wb["Corrections"].values)[1:]}
    maximum=int(wb["Policy"]["B2"].value)
    cities=str(wb["Policy"]["B3"].value).split(",")
    eligible={}
    for name,price,qty,bj,sz,status in list(wb["Orders"].values)[1:]:
        if status!="active" or ("Beijing" in cities and not bj) or ("Shenzhen" in cities and not sz):continue
        total=int(price)*corr.get(str(name),int(qty))
        if total<=maximum:eligible[str(name)]=total
    if not eligible:raise ValueError("no eligible orders")
    winner=min(eligible,key=lambda name:(-eligible[name],name))
    return tuple(map(str,(sum(eligible.values()),len(eligible),winner,eligible[winner])))

def report(answer):
    return f"Total: {answer[0]}; Count: {answer[1]}; Winner: {answer[2]}; Amount: {answer[3]}."

def report_valid(text,answer):
    for label,value in zip(("Total","Count","Winner","Amount"),answer):
        mentions=re.findall(r"(?i)\b"+label+r"\s*:\s*([A-Za-z0-9-]+)",text)
        if not mentions or any(m!=value for m in mentions):return False
    return True

def specs():
    s={"type":"string"}
    def tool(name,props=None):
        return {"type":"function","function":{"name":name,"description":name.replace("_"," "),"parameters":{"type":"object","properties":props or {},"required":list((props or {}).keys()),"additionalProperties":False}}}
    return [tool("list_sheets"),tool("read_sheet",{"sheet":s}),tool("write_summary",{"cell":s,"value":s}),tool("save_workbook"),tool("write_report",{"text":s}),tool("receive_update"),tool("finish")]

class Environment:
    def __init__(self,case,folder):
        self.case=case;self.root=Path(folder);self.path=self.root/case["input"]["workbook"]
        self.wb=load_workbook(self.path);self.calls=[];self.invalid_actions=0
        self.finished=False;self.transitioned=False;self.initial_ok=False
        self.saved_after=False;self.old_report=None
    def call(self,name,args):
        if self.finished:raise ValueError("already finished")
        if name=="list_sheets" and args=={}:out={"sheets":self.wb.sheetnames}
        elif name=="read_sheet" and set(args)=={"sheet"} and args["sheet"] in self.wb.sheetnames:
            out={"rows":[list(row) for row in self.wb[args["sheet"]].values]}
        elif name=="write_summary" and set(args)=={"cell","value"} and args["cell"] in ("B2","B3","B4","B5") and isinstance(args["value"],str):
            self.wb["Summary"][args["cell"]]=args["value"];out={"ok":True}
        elif name=="save_workbook" and args=={}:
            self.wb.save(self.path);out={"saved":True}
            if self.transitioned:self.saved_after=True
        elif name=="write_report" and set(args)=={"text"} and isinstance(args["text"],str):
            (self.root/"report.txt").write_text(args["text"],encoding="utf-8");out={"saved":True}
        elif name=="receive_update" and args=={} and not self.transitioned:
            disk=load_workbook(self.path,read_only=True,data_only=True)
            try:
                first=oracle(disk)
                actual=tuple(str(disk["Summary"][c].value or "") for c in ("B2","B3","B4","B5"))
            finally:disk.close()
            p=self.root/"report.txt"
            old=p.read_text(encoding="utf-8") if p.exists() else ""
            self.initial_ok=(actual==first and report_valid(old,first))
            if not self.initial_ok:raise ValueError("first-stage deliverables incorrect or unsaved")
            self.old_report=old;self.transitioned=True
            self.wb["Policy"]["B2"]=8000;self.wb["Policy"]["B3"]="Beijing,Shenzhen"
            quantities=((4,2,1),(1,3,1),(1,1,3))[self.case["scenario"]]
            for row,qty in enumerate(quantities,2):self.wb["Corrections"].cell(row,2,qty)
            self.wb.save(self.path)
            out={"max_per_order":8000,"cities":["Beijing","Shenzhen"],"changed_sheets":["Policy","Corrections"]}
        elif name=="finish" and args=={} and self.transitioned:
            self.finished=True;out={"ok":True}
        else:raise ValueError("invalid tool call or stage")
        self.calls.append({"tool":name,"args":dict(args),"response":out})
        return out
    def score(self):
        disk=load_workbook(self.path,read_only=True,data_only=True)
        try:
            correct=oracle(disk)
            actual=tuple(str(disk["Summary"][c].value or "") for c in ("B2","B3","B4","B5"))
            policy=(str(disk["Policy"]["B2"].value)=="8000" and
                    str(disk["Policy"]["B3"].value)=="Beijing,Shenzhen")
            expected_q=((4,2,1),(1,3,1),(1,1,3))[self.case["scenario"]]
            quantities=tuple(disk["Corrections"].cell(i,2).value for i in range(2,5))
            policy=policy and quantities==expected_q
        finally:disk.close()
        p=self.root/"report.txt";text=p.read_text(encoding="utf-8") if p.is_file() else ""
        checks={"initial_delivered":self.initial_ok,"update_received":self.transitioned,
                "requirements":policy,"result":actual==correct,
                "report":report_valid(text,correct) and text!=self.old_report,
                "revised_saved":self.saved_after,"finished":self.finished}
        return {"success":all(checks.values()) and not self.invalid_actions,
                "checks":checks,"expected":list(correct),"actual":list(actual),
                "invalid_actions":self.invalid_actions,"tool_calls":len(self.calls)}
    def close(self):self.wb.close()

def reference(case,folder):
    env=Environment(case,folder)
    try:
        for stage in range(2):
            for sheet in ("Orders","Corrections","Policy"):env.call("read_sheet",{"sheet":sheet})
            answer=oracle(env.wb)
            for cell,value in zip(("B2","B3","B4","B5"),answer):
                env.call("write_summary",{"cell":cell,"value":value})
            env.call("save_workbook",{})
            env.call("write_report",{"text":report(answer)})
            if stage==0:env.call("receive_update",{})
        env.call("finish",{})
        return env.score()
    finally:env.close()
