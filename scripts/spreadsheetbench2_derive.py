"""Produce six *candidate* tasks from one immutable real XLSX, with private checks.

This is a controlled research generator, not a validated official benchmark.
"""
import argparse,hashlib,json,shutil,re
from pathlib import Path
from openpyxl import load_workbook

SHEET="Ex 5 - M&A"
ROWS={42:(36,41),52:(46,51),62:(56,61),72:(66,71)}
COLS="DEFG"
TARGETS={f"{c}{r}":f"=+SUM({c}{lo}:{c}{hi})"
         for r,(lo,hi) in ROWS.items() for c in COLS}
CASES=[
 ("locate","Identify the cells in 'Ex 5 - M&A' whose totals incorrectly accumulate prior-period subtotals. Return JSON with a cells array; do not edit the workbook.","cell_set"),
 ("repair","Correct the double-counting formulas in the M&A subtotal rows. Return an edited XLSX; leave unrelated cell contents untouched.","minimal_repair"),
 ("row42","Correct only the first affected subtotal row (row 42) in 'Ex 5 - M&A'. Return an edited XLSX; do not change other cells.","minimal_repair"),
 ("row52","Correct only the affected subtotal row 52 in 'Ex 5 - M&A'. Return an edited XLSX; do not change other cells.","minimal_repair"),
 ("row62","Correct only the affected subtotal row 62 in 'Ex 5 - M&A'. Return an edited XLSX; do not change other cells.","minimal_repair"),
 ("row72","Correct only the affected subtotal row 72 in 'Ex 5 - M&A'. Return an edited XLSX; do not change other cells.","minimal_repair"),
]
def digest(path):
 h=hashlib.sha256()
 with path.open("rb") as f:
  for chunk in iter(lambda:f.read(1024*1024),b""):h.update(chunk)
 return h.hexdigest()
def normalized(value):
 if value is None or isinstance(value,(str,int,float,bool)):return value
 return {"type":type(value).__name__,"text":getattr(value,"text",None),"ref":getattr(value,"ref",None),"dt2D":getattr(value,"dt2D",None)}
def values(path):
 w=load_workbook(path,read_only=False,data_only=False)
 try:
  return {(s.title,c.coordinate):normalized(c.value)
    for s in w for c in s._cells.values()}
 finally:w.close()
def formula_matches(actual, expected):
 # Only normalize a mechanically verifiable SUM of the exact same contiguous
 # cells. Never claim general mathematical/formula equivalence.
 if actual==expected:return True
 if not isinstance(actual,str) or not isinstance(expected,str):return False
 m=re.fullmatch(r"=\\+?SUM\\(([A-Z]+)([0-9]+):([A-Z]+)([0-9]+)\\)",expected,re.I)
 if not m or m.group(1).upper()!=m.group(3).upper():return False
 col=m.group(1).upper();lo=int(m.group(2));hi=int(m.group(4))
 terms=[f"{col}{row}" for row in range(lo,hi+1)]
 expression=actual.lstrip("=+").upper().replace("$","").replace(" ","")
 return expression in ("+".join(terms),"+".join(reversed(terms)))

def structure(path):
 w=load_workbook(path,read_only=False,data_only=False,keep_links=True)
 try:
  return {"sheets":w.sheetnames,
          "merged":{s.title:sorted(str(r) for r in s.merged_cells.ranges) for s in w},
          "defined_names":sorted(w.defined_names),
          "sheet_states":{s.title:s.sheet_state for s in w}}
 finally:w.close()

def validate(case,submission):
 case=Path(case);meta=json.loads((case/"private"/"oracle.json").read_text(encoding="utf-8"))
 if meta["kind"]=="cell_set":
  answer=json.loads(Path(submission).read_text(encoding="utf-8"))
  actual=set(answer.get("cells",[]))
  expected=set(meta["expected_cells"])
  return {"pass":actual==expected,"missing":sorted(expected-actual),"extra":sorted(actual-expected)}
 input_file=next((case/"input").glob("*.xlsx"))
 original=values(input_file);edited=values(Path(submission))
 target={(SHEET,k):v for k,v in meta["expected_formulas"].items()}
 wrong={f"{s}!{cell}":{"expected":val,"actual":edited.get((s,cell))}
        for (s,cell),val in target.items() if not formula_matches(edited.get((s,cell)),val)}
 extra=[f"{s}!{cell}" for (s,cell) in original.keys()|edited.keys()
        if (s,cell) not in target and original.get((s,cell))!=edited.get((s,cell))]
 structural_differences={}
 before_structure=structure(input_file);after_structure=structure(Path(submission))
 for key in before_structure:
  if before_structure[key]!=after_structure[key]:
   structural_differences[key]={"original":before_structure[key],"submitted":after_structure[key]}
 return {"pass":not wrong and not extra and not structural_differences,
         "wrong":wrong,"unexpected_edits":sorted(extra)[:50],
         "unexpected_count":len(extra),"structure_differences":structural_differences,
         "note":"Exact formulas plus verified SUM expansions; limited workbook structure checks. Not a numerical recalculation or full OOXML audit."}
def generate(source,out):
 source=Path(source).resolve();out=Path(out).resolve()
 if out.exists():raise FileExistsError(out)
 wb=load_workbook(source,read_only=False,data_only=False)
 try:
  s=wb[SHEET]
  for cell in TARGETS:
   if not isinstance(s[cell].value,str) or "+SUM(" not in s[cell].value:
    raise ValueError(f"Unexpected original formula at {cell}: {s[cell].value!r}")
 finally:wb.close()
 out.mkdir(parents=True)
 index=[]
 for slug,instruction,kind in CASES:
  folder=out/slug
  for sub in ("input","output","private"):(folder/sub).mkdir(parents=True)
  shutil.copy2(source,folder/"input"/source.name)
  subset=TARGETS if slug in ("locate","repair") else {k:v for k,v in TARGETS.items() if k.endswith(slug[3:])}
  oracle={"kind":kind,"source_sha256":digest(source),"expected_cells":sorted(subset),
          "expected_formulas":subset,"status":"candidate_requires_review",
          "source":"SpreadsheetBench 2 Debugging 01_01; Golden-confirmed formulas"}
  (folder/"instruction.md").write_text(instruction,encoding="utf-8")
  (folder/"private"/"oracle.json").write_text(json.dumps(oracle,indent=2),encoding="utf-8")
  index.append({"case":slug,"kind":kind,"workspace":str(folder)})
 (out/"index.json").write_text(json.dumps(index,indent=2),encoding="utf-8")
 return index
def selftest(out):
 results=[]
 for case in sorted(Path(out).iterdir()):
  if not case.is_dir():continue
  src=next((case/"input").glob("*.xlsx"))
  oracle=json.loads((case/"private"/"oracle.json").read_text(encoding="utf-8"))
  if oracle["kind"]=="cell_set":
   good=case/"private"/"correct.json";bad=case/"private"/"incorrect.json"
   good.write_text(json.dumps({"cells":oracle["expected_cells"]}),encoding="utf-8")
   bad.write_text('{"cells":[]}',encoding="utf-8")
  else:
   bad=src
   good=case/"private"/"correct.xlsx"
   wb=load_workbook(src)
   for addr,formula in oracle["expected_formulas"].items():wb[SHEET][addr]=formula
   wb.save(good);wb.close()
  good_result=validate(case,good);bad_result=validate(case,bad)
  results.append({"case":case.name,"positive_pass":good_result["pass"],
                  "negative_rejected":not bad_result["pass"],"oracle_status":oracle["status"]})
 return results
def main():
 p=argparse.ArgumentParser()
 p.add_argument("action",choices=["generate","selftest","grade"])
 p.add_argument("--source",type=Path,default=Path("runs/ssb2-debugging-baseline/01_Debugging_01_01/input/Double Counting_input.xlsx"))
 p.add_argument("--out",type=Path,default=Path("runs/ssb2-derived-01"))
 p.add_argument("--case",type=Path);p.add_argument("--submission",type=Path)
 a=p.parse_args()
 if a.action=="generate":result=generate(a.source,a.out)
 elif a.action=="selftest":result=selftest(a.out)
 else:
  if not a.case or not a.submission:p.error("grade requires --case and --submission")
  result=validate(a.case,a.submission)
 print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=="__main__":main()
