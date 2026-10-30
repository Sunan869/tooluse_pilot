"""Trusted-side diagnostic: compare official SpreadsheetBench input and golden XLSX.

Run ONLY on the host, never mount private/ or this script in the agent.
A golden diff is evidence of changes, not a substitute for official grading.
"""
import argparse
import json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

def normalize(value):
    if value is None or isinstance(value,(str,int,float,bool)):
        return value
    return {"kind":type(value).__name__,
            "text":getattr(value,"text",None),
            "ref":getattr(value,"ref",None),
            "dt2D":getattr(value,"dt2D",None)}

def style_signature(cell):
    if cell is None:
        return None
    # Style IDs are workbook-local indexes, not comparable across two files.
    # Compare the underlying actual style objects (fonts, fills, etc.).
    return {
        "font": str(cell.font),
        "fill": str(cell.fill),
        "border": str(cell.border),
        "alignment": str(cell.alignment),
        "number_format": cell.number_format,
        "protection": str(cell.protection),
    }

def compare(case, max_examples=60):
    case=Path(case)
    inputs=list((case/"input").glob("*.xlsx"))
    golds=list((case/"private").glob("*.xlsx"))
    if len(inputs)!=1 or len(golds)!=1:
        raise ValueError(f"Expected one input and one golden workbook; found {len(inputs)} and {len(golds)}")
    a=load_workbook(inputs[0],read_only=False,data_only=False,keep_links=True)
    b=load_workbook(golds[0],read_only=False,data_only=False,keep_links=True)
    changes=[]
    by_sheet=Counter()
    kinds=Counter()
    value_categories=Counter()
    value_by_sheet=Counter()
    non_table_examples=[]
    value_examples=[]
    style_examples=[]
    try:
        for name in sorted(set(a.sheetnames)|set(b.sheetnames)):
            if name not in a or name not in b:
                by_sheet[name]+=1
                kinds["sheet_added_or_removed"]+=1
                changes.append({"sheet":name,"kind":"sheet_added_or_removed"})
                continue
            old,new=a[name],b[name]
            coords=set(old._cells)|set(new._cells)
            for row,col in sorted(coords):
                x=old._cells.get((row,col));y=new._cells.get((row,col))
                before=normalize(x.value) if x else None
                after=normalize(y.value) if y else None
                style_before=style_signature(x)
                style_after=style_signature(y)
                value_changed=before!=after
                style_changed=style_before!=style_after
                if not value_changed and not style_changed: continue
                by_sheet[name]+=1
                kind="value_or_formula" if value_changed else "style_only"
                kinds[kind]+=1
                if value_changed:
                    table_related = (
                        (isinstance(before,dict) and before.get("kind")=="DataTableFormula")
                        or (isinstance(after,dict) and after.get("kind")=="DataTableFormula")
                        or (isinstance(before,str) and before.startswith("=TABLE("))
                        or (isinstance(after,str) and after.startswith("=TABLE("))
                    )
                    value_category = "excel_data_table_conversion" if table_related else "other_value_or_formula"
                    value_categories[value_category]+=1
                    value_by_sheet[name]+=1
                    if not table_related and len(non_table_examples)<max_examples:
                        non_table_examples.append({"sheet":name,"cell":old.cell(row,col).coordinate,
                                                   "before":before,"after":after})
                target = value_examples if value_changed else style_examples
                if len(target)<max_examples:
                    target.append({"sheet":name,"cell":f"{old.cell(row,col).coordinate}",
                                   "kind":kind,"before":before,"after":after,
                                   "style_changed":style_changed})
        return {"case":str(case),"input":inputs[0].name,"golden":golds[0].name,
                "total_changed_cells":sum(by_sheet.values()),
                "by_sheet":dict(by_sheet),"change_kinds":dict(kinds),
                "value_change_categories":dict(value_categories),
                "value_changes_by_sheet":dict(value_by_sheet),
                "non_data_table_examples":non_table_examples,
                "examples":(value_examples + style_examples)[:max_examples],
                "value_or_formula_examples":value_examples,
                "style_only_examples":style_examples[:max(0,max_examples-len(value_examples))],
                "examples_truncated":sum(by_sheet.values())>max_examples,
                "notice":"Excel data-table conversions and widespread style differences may reflect workbook conversion. This is not an official score or proof of necessary fixes."}
    finally:
        a.close();b.close()

def main():
    p=argparse.ArgumentParser()
    p.add_argument("case",type=Path)
    p.add_argument("--out",type=Path)
    p.add_argument("--max-examples",type=int,default=60)
    a=p.parse_args()
    result=compare(a.case,max(0,a.max_examples))
    if a.out:
        a.out.parent.mkdir(parents=True,exist_ok=True)
        a.out.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
    print(json.dumps(result,ensure_ascii=False,indent=2,default=str))

if __name__=="__main__":main()
