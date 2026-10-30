"""Read-only XLSX difference audit, not the official benchmark grader."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook

def describe(value):
    if value is None or isinstance(value,(str,int,float,bool)): return value
    return {"type":type(value).__name__,"text":getattr(value,"text",None),"ref":getattr(value,"ref",None)}

def audit(case):
    case=Path(case)
    original=list((case/"input").glob("*.xlsx"))
    edited=list((case/"output").glob("*.xlsx"))
    if len(original)!=1 or len(edited)!=1:
        return {"status":"no_comparable_artifact","input_count":len(original),"output_count":len(edited)}
    a=load_workbook(original[0],read_only=False,data_only=False)
    b=load_workbook(edited[0],read_only=False,data_only=False)
    changes=[]
    try:
        for name in sorted(set(a.sheetnames)|set(b.sheetnames)):
            if name not in a or name not in b:
                changes.append({"sheet":name,"change":"sheet_added_or_removed"})
                continue
            x,y=a[name],b[name]
            keys=set(x._cells)|set(y._cells)
            for row,col in sorted(keys):
                before=x.cell(row,col);after=y.cell(row,col)
                old=describe(before.value);new=describe(after.value)
                if old!=new or before.style_id!=after.style_id:
                    changes.append({"sheet":name,"cell":before.coordinate,
                                    "old":old,"new":new,
                                    "style_changed":before.style_id!=after.style_id})
        return {"status":"compared","input":original[0].name,"output":edited[0].name,
                "sheet_counts":[len(a.sheetnames),len(b.sheetnames)],
                "changed_cells":len(changes),"changes":changes[:200],
                "truncated":len(changes)>200,
                "warning":"Formula correctness, recalculation and official score NOT checked"}
    finally:
        a.close();b.close()

def main():
    p=argparse.ArgumentParser()
    p.add_argument("case")
    p.add_argument("--out",type=Path)
    args=p.parse_args()
    result=audit(args.case)
    if args.out:
        args.out.parent.mkdir(parents=True,exist_ok=True)
        args.out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({k:v for k,v in result.items() if k!="changes"},ensure_ascii=False,indent=2))
if __name__=="__main__":main()
