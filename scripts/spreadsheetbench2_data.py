"""Download, inspect and shortlist real SpreadsheetBench 2 workbooks.

Trusted-side only. Never mount this module's dataset root or golden files in an
agent container; stage only an individual input workbook and its instruction.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path

DATASET = "KAKA22/SpreadsheetBench-v2"
ARCHIVE = "spreadsheetbench-v2-09-10.zip"
CATEGORIES = ("Debugging", "Financial_Model", "Template")
ROOT = Path("datasets/spreadsheetbench2")


def safe_extract(archive: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    base = destination.resolve()
    with zipfile.ZipFile(archive) as zf:
        for info in zf.infolist():
            name = info.filename.replace("\\", "/")
            target = (base / name).resolve()
            if not target.is_relative_to(base) or info.is_dir() and name.startswith("/"):
                raise ValueError(f"Unsafe archive member: {name}")
            if info.file_size > 512 * 1024 * 1024:
                raise ValueError(f"Excessively large archive member: {name}")
            if info.is_dir():
                target.mkdir(parents=True, exist_ok=True)
            else:
                target.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(info) as source, target.open("wb") as output:
                    shutil.copyfileobj(source, output)


def download(root: Path, archive: str = ARCHIVE) -> None:
    root.mkdir(parents=True, exist_ok=True)
    cmd = [sys.executable, "-m", "huggingface_hub.commands.huggingface_cli"]
    # Use the supported public Python API instead of relying on a CLI entrypoint.
    try:
        from huggingface_hub import hf_hub_download
    except ImportError as exc:
        raise SystemExit("Install: python -m pip install huggingface_hub") from exc
    path = Path(hf_hub_download(repo_id=DATASET, repo_type="dataset",
                                filename=archive, local_dir=root / "downloads"))
    print(f"Downloaded: {path} ({path.stat().st_size:,} bytes)")
    safe_extract(path, root / "raw")
    print(f"Extracted to: {root / 'raw'}")


def summarize_workbook(path: Path, max_rows: int = 2000, max_cols: int = 80) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=True)
    sheets = len(wb.worksheets)
    cells = formulas = cross_sheet = errors = 0
    functions = Counter()
    try:
        for sheet in wb:
            for row in sheet.iter_rows(max_row=min(sheet.max_row or 0, max_rows),\n                                       max_col=min(sheet.max_column or 0, max_cols)):
                for cell in row:
                    if cell.value is None:
                        continue
                    cells += 1
                    if cell.data_type == "e":
                        errors += 1
                    if cell.data_type == "f":
                        formulas += 1
                        expression = str(cell.value)
                        cross_sheet += "!" in expression
                        head = expression.lstrip("=").split("(", 1)[0].upper()
                        if "(" in expression and head.isidentifier():
                            functions[head] += 1
    finally:
        wb.close()
    return {"sheets": sheets, "occupied_cells": cells, "formulas": formulas,
            "cross_sheet_formulas": cross_sheet, "error_cells": errors,
            "top_functions": dict(functions.most_common(8)),\n            "sample_limit": {"max_rows_per_sheet": max_rows, "max_cols_per_sheet": max_cols},\n            "partial_scan": any((ws.max_row or 0) > max_rows or\n                                (ws.max_column or 0) > max_cols for ws in wb.worksheets)}


def catalog(root: Path, limit: int, output: Path) -> list:
    raw = root / "raw"
    if not raw.exists():
        raise SystemExit(f"Missing {raw}; run download or manually extract first")
    records = []
    for path in sorted(raw.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        category = next((name for name in CATEGORIES if name.lower() in
                         [part.lower() for part in path.parts]), None)
        if category is None:
            continue
        try:
            stat = summarize_workbook(path)
            records.append({"path": path.relative_to(raw).as_posix(),
                            "category": category, "bytes": path.stat().st_size,
                            **stat})
        except Exception as exc:
            records.append({"path": path.relative_to(raw).as_posix(),
                            "category": category, "error": str(exc)})
        if limit and len(records) >= limit:
            break
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Inspected {len(records)} workbooks; manifest: {output}", flush=True)
    return records


def shortlist(manifest: Path, out: Path, count: int) -> None:
    records = json.loads(manifest.read_text(encoding="utf-8"))
    eligible = [r for r in records if "error" not in r and r["formulas"] > 0]
    # Rank formula diversity and cross-sheet references, not file size alone.
    eligible.sort(key=lambda r: (r["cross_sheet_formulas"] > 0,
                  len(r["top_functions"]), r["cross_sheet_formulas"],
                  r["formulas"], r["sheets"]), reverse=True)
    selected = eligible[:count]
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({"method": "formula_and_dependency_preselection",
        "note": "Structural candidates only; not verified benchmark cases or ground truth.",
        "cases": selected}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Selected {len(selected)} candidates: {out}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=("download", "catalog", "shortlist"))
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--archive", default=ARCHIVE)
    parser.add_argument("--limit", type=int, default=0,
                        help="Max workbooks to inspect (0=all)")
    parser.add_argument("--count", type=int, default=20)
    args = parser.parse_args()
    if args.action == "download":
        download(args.root, args.archive)
    elif args.action == "catalog":
        catalog(args.root, args.limit, args.root / "manifest.json")
    else:
        shortlist(args.root / "manifest.json", args.root / "shortlist.json", args.count)


if __name__ == "__main__":
    main()
