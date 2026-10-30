"""Independent negative-mutation checks against the derived-task oracle.

Uses trusted fixture only; never calls a model, never modifies original inputs.
"""
import argparse
import json
import sys
import tempfile
from pathlib import Path
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parent))
from spreadsheetbench2_derive import validate, SHEET

def test_case(case):
 case=Path(case)
 if not (case/"private"/"oracle.json").exists():raise FileNotFoundError(case)
 oracle=json.loads((case/"private"/"oracle.json").read_text(encoding="utf-8"))
 checks={}
 if oracle["kind"]=="cell_set":
  expected=oracle["expected_cells"]
  with tempfile.TemporaryDirectory() as td:
   folder=Path(td)
   for name,cells,want in (
    ("correct",expected,True),
    ("missing",expected[:-1],False),
    ("extra",expected+["A1"],False)):
    file=folder/f"{name}.json"
    file.write_text(json.dumps({"cells":cells}),encoding="utf-8")
    checks[name]={"actual":validate(case,file)["pass"],"expected":want}
 else:
  original=next((case/"input").glob("*.xlsx"))
  target=oracle["expected_formulas"]
  with tempfile.TemporaryDirectory() as td:
   tmp=Path(td)
   def make(name,mutator):
    book=load_workbook(original)
    for address,formula in target.items():book[SHEET][address]=formula
    mutator(book)
    output=tmp/f"{name}.xlsx"
    book.save(output);book.close()
    return output
   address=next(iter(target))
   def noop(book):pass
   def omitted(book):
    source=load_workbook(original)
    book[SHEET][address]=source[SHEET][address].value
    source.close()
   def wrong_reference(book):
    book[SHEET][address]="=+SUM(A1:A2)"
   def unauthorized(book):
    book[SHEET]["B2"]="UNAUTHORIZED_CHANGE"
   def missing_sheet(book):
    book.remove(book.worksheets[-1])
   variants=[("correct",noop,True),("omitted",omitted,False),
             ("wrong_reference",wrong_reference,False),
             ("unauthorized",unauthorized,False),
             ("missing_sheet",missing_sheet,False)]
   for name,mutation,want in variants:
    result=validate(case,make(name,mutation))
    checks[name]={"actual":result["pass"],"expected":want}
 return {"case":case.name,"passed":all(v["actual"]==v["expected"] for v in checks.values()),"checks":checks}

def main():
 p=argparse.ArgumentParser()
 p.add_argument("--root",type=Path,default=Path("runs/ssb2-derived-01"))
 a=p.parse_args()
 results=[test_case(path) for path in sorted(a.root.iterdir()) if path.is_dir()]
 report={"passed":all(item["passed"] for item in results),"cases":results}
 print(json.dumps(report,ensure_ascii=False,indent=2))
 if not report["passed"]:raise SystemExit(1)
if __name__=="__main__":main()
